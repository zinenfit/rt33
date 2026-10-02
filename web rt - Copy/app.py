from flask import (
    Flask,
    render_template,
    request,
    redirect,
    send_file,
    make_response,
    session
)

import sqlite3

from werkzeug.security import (
    generate_password_hash,
    check_password_hash
)

from functools import wraps
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook

from itsdangerous import (
    URLSafeTimedSerializer,
    BadSignature,
    SignatureExpired
)

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import (
    WD_TABLE_ALIGNMENT,
    WD_CELL_VERTICAL_ALIGNMENT
)


# ============================================================
# APP
# ============================================================

app = Flask(__name__)

app.secret_key = "ganti-dengan-kunci-rahasia-sendiri"

DATABASE = "rt.db"

ADMIN_USERNAME = "adminrt"
ADMIN_PASSWORD = "AdminRT2026"

COOKIE_MAX_AGE = 60 * 60 * 24 * 7

ADMIN_COOKIE = "rt_admin_auth"
USER_COOKIE = "rt_user_auth"

serializer = URLSafeTimedSerializer(app.secret_key)


# ============================================================
# DATABASE
# ============================================================

def get_db():
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():

    conn = get_db()
    cursor = conn.cursor()

    # --------------------------------------------------------
    # USERS
    # --------------------------------------------------------

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL,
            role TEXT NOT NULL
        )
    """)

    # --------------------------------------------------------
    # WARGA
    # --------------------------------------------------------

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS warga (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nama TEXT NOT NULL,
            nik TEXT NOT NULL,
            alamat TEXT NOT NULL,
            user_id INTEGER NOT NULL,
            FOREIGN KEY (user_id) REFERENCES users(id)
        )
    """)

    # --------------------------------------------------------
    # PENGUMUMAN
    # --------------------------------------------------------

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS pengumuman (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            judul TEXT NOT NULL,
            isi TEXT NOT NULL,
            tanggal TEXT NOT NULL
        )
    """)

    # --------------------------------------------------------
    # PENGAJUAN SURAT
    # --------------------------------------------------------

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS pengajuan_surat (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            jenis_surat TEXT NOT NULL,
            keperluan TEXT NOT NULL,
            keterangan TEXT,
            tanggal_pengajuan TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'Menunggu',
            alasan_penolakan TEXT,
            nomor_surat TEXT,
            FOREIGN KEY (user_id) REFERENCES users(id)
        )
    """)

    # --------------------------------------------------------
    # BUAT ADMIN JIKA BELUM ADA
    # --------------------------------------------------------

    admin = cursor.execute("""
        SELECT id
        FROM users
        WHERE role = 'admin'
        LIMIT 1
    """).fetchone()

    if admin is None:

        password_hash = generate_password_hash(
            ADMIN_PASSWORD
        )

        cursor.execute("""
            INSERT INTO users
            (
                username,
                password,
                role
            )
            VALUES (?, ?, ?)
        """, (
            ADMIN_USERNAME,
            password_hash,
            "admin"
        ))

    conn.commit()
    conn.close()


# ============================================================
# AUTH TOKEN
# ============================================================

def create_auth_token(user_id, username, role):

    return serializer.dumps({
        "user_id": user_id,
        "username": username,
        "role": role
    })


def read_auth_token(cookie_name, expected_role):

    token = request.cookies.get(cookie_name)

    if not token:
        return None

    try:

        data = serializer.loads(
            token,
            max_age=COOKIE_MAX_AGE
        )

    except (
        BadSignature,
        SignatureExpired
    ):

        return None

    if data.get("role") != expected_role:
        return None

    conn = get_db()

    user = conn.execute("""
        SELECT
            id,
            username,
            role
        FROM users
        WHERE id = ?
        AND role = ?
    """, (
        data.get("user_id"),
        expected_role
    )).fetchone()

    conn.close()

    if user is None:
        return None

    return {
        "user_id": user["id"],
        "username": user["username"],
        "role": user["role"]
    }


# ============================================================
# ADMIN AUTH
# ============================================================

def get_admin():

    return read_auth_token(
        ADMIN_COOKIE,
        "admin"
    )


# ============================================================
# USER AUTH
# ============================================================

def get_user():

    return read_auth_token(
        USER_COOKIE,
        "user"
    )


# ============================================================
# CURRENT USER
#
# Dipakai hanya untuk halaman umum yang memang menerima
# Admin maupun User.
#
# Untuk /admin gunakan get_admin()
# Untuk /user gunakan get_user()
# ============================================================

def get_current_user():

    admin = get_admin()

    if admin:
        return admin

    user = get_user()

    if user:
        return user

    return None


# ============================================================
# SESSION
#
# Session hanya untuk informasi tampilan.
# Hak akses TIDAK ditentukan oleh session.
# Hak akses ditentukan oleh cookie Admin/User.
# ============================================================

def set_session(user):

    session.clear()

    session["user_id"] = user["user_id"]
    session["username"] = user["username"]
    session["role"] = user["role"]


# ============================================================
# ADMIN REQUIRED
# ============================================================

def admin_required(function):

    @wraps(function)
    def decorated_function(*args, **kwargs):

        admin = get_admin()

        if admin is None:

            session.clear()

            return redirect("/")

        set_session(admin)

        return function(*args, **kwargs)

    return decorated_function


# ============================================================
# USER REQUIRED
# ============================================================

def user_required(function):

    @wraps(function)
    def decorated_function(*args, **kwargs):

        user = get_user()

        if user is None:

            session.clear()

            return redirect("/")

        set_session(user)

        return function(*args, **kwargs)

    return decorated_function


# ============================================================
# LOGIN REQUIRED
# ============================================================

def login_required(function):

    @wraps(function)
    def decorated_function(*args, **kwargs):

        current = get_current_user()

        if current is None:

            session.clear()

            return redirect("/")

        set_session(current)

        return function(*args, **kwargs)

    return decorated_function


# ============================================================
# LOGIN
# ============================================================

@app.route("/", methods=["GET", "POST"])
def login():

    # --------------------------------------------------------
    # GET
    # --------------------------------------------------------

    if request.method == "GET":

        admin = get_admin()
        user = get_user()

        # Jika hanya Admin yang login
        if admin and not user:

            set_session(admin)

            return redirect("/admin")

        # Jika hanya User yang login
        if user and not admin:

            set_session(user)

            return redirect("/user")

        # Jika kedua cookie ada sekaligus,
        # jangan memilih salah satu secara otomatis.
        #
        # Tampilkan login supaya user dapat masuk
        # dengan akun yang diinginkan.
        if admin and user:

            session.clear()

            return render_template("login.html")

        session.clear()

        return render_template("login.html")

    # --------------------------------------------------------
    # POST LOGIN
    # --------------------------------------------------------

    username = request.form.get(
        "username",
        ""
    ).strip()

    password = request.form.get(
        "password",
        ""
    )

    if username == "" or password == "":

        return render_template(
            "login.html",
            error="Username dan password harus diisi."
        )

    conn = get_db()

    user = conn.execute("""
        SELECT *
        FROM users
        WHERE username = ?
    """, (
        username,
    )).fetchone()

    conn.close()

    if user is None:

        return render_template(
            "login.html",
            error="Username atau password salah."
        )

    if not check_password_hash(
        user["password"],
        password
    ):

        return render_template(
            "login.html",
            error="Username atau password salah."
        )

    # --------------------------------------------------------
    # LOGIN ADMIN
    # --------------------------------------------------------

    if user["role"] == "admin":

        token = create_auth_token(
            user["id"],
            user["username"],
            "admin"
        )

        session.clear()

        set_session({
            "user_id": user["id"],
            "username": user["username"],
            "role": "admin"
        })

        response = make_response(
            redirect("/admin")
        )

        response.set_cookie(
            ADMIN_COOKIE,
            token,
            max_age=COOKIE_MAX_AGE,
            httponly=True,
            samesite="Lax"
        )

        return response

    # --------------------------------------------------------
    # LOGIN USER
    # --------------------------------------------------------

    token = create_auth_token(
        user["id"],
        user["username"],
        "user"
    )

    session.clear()

    set_session({
        "user_id": user["id"],
        "username": user["username"],
        "role": "user"
    })

    response = make_response(
        redirect("/user")
    )

    # Hapus cookie Admin agar tidak bentrok
    response.delete_cookie(ADMIN_COOKIE)

    response.set_cookie(
        USER_COOKIE,
        token,
        max_age=COOKIE_MAX_AGE,
        httponly=True,
        samesite="Lax"
    )

    return response


# ============================================================
# ADMIN DASHBOARD
# ============================================================

@app.route("/admin")
@admin_required
def admin():

    admin_data = get_admin()

    conn = get_db()

    # --------------------------------------------------------
    # WARGA
    # --------------------------------------------------------

    warga = conn.execute("""
        SELECT
            warga.id,
            warga.nama,
            warga.nik,
            warga.alamat,
            warga.user_id,
            users.username
        FROM warga
        LEFT JOIN users
        ON warga.user_id = users.id
        ORDER BY warga.id DESC
    """).fetchall()

    # --------------------------------------------------------
    # USERS
    # --------------------------------------------------------

    users = conn.execute("""
        SELECT
            id,
            username,
            role
        FROM users
        WHERE role = 'user'
        ORDER BY id DESC
    """).fetchall()

    # --------------------------------------------------------
    # PENGUMUMAN
    # --------------------------------------------------------

    pengumuman = conn.execute("""
        SELECT *
        FROM pengumuman
        ORDER BY id DESC
    """).fetchall()

    # --------------------------------------------------------
    # SURAT
    # --------------------------------------------------------

    pengajuan = conn.execute("""
        SELECT
            pengajuan_surat.*,
            users.username,
            warga.nama,
            warga.nik,
            warga.alamat
        FROM pengajuan_surat

        LEFT JOIN users
        ON pengajuan_surat.user_id = users.id

        LEFT JOIN warga
        ON pengajuan_surat.user_id = warga.user_id

        ORDER BY pengajuan_surat.id DESC
    """).fetchall()

    conn.close()

    return render_template(
        "admin.html",
        warga=warga,
        users=users,
        pengumuman=pengumuman,
        pengajuan=pengajuan,
        username=admin_data["username"]
    )


# ============================================================
# USER DASHBOARD
# ============================================================

@app.route("/user")
@user_required
def user():

    user_data = get_user()

    conn = get_db()

    # --------------------------------------------------------
    # DATA WARGA
    # --------------------------------------------------------

    warga = conn.execute("""
        SELECT
            id,
            nama,
            nik,
            alamat,
            user_id
        FROM warga
        WHERE user_id = ?
        ORDER BY id DESC
    """, (
        user_data["user_id"],
    )).fetchall()

    # --------------------------------------------------------
    # PENGUMUMAN
    # --------------------------------------------------------

    pengumuman = conn.execute("""
        SELECT *
        FROM pengumuman
        ORDER BY id DESC
    """).fetchall()

    # --------------------------------------------------------
    # SURAT
    # --------------------------------------------------------

    pengajuan = conn.execute("""
        SELECT *
        FROM pengajuan_surat
        WHERE user_id = ?
        ORDER BY id DESC
    """, (
        user_data["user_id"],
    )).fetchall()

    conn.close()

    return render_template(
        "user.html",
        warga=warga,
        pengumuman=pengumuman,
        pengajuan=pengajuan,
        username=user_data["username"]
    )


# ============================================================
# REGISTER AKUN
#
# PENTING:
# Membuat akun saja.
# TIDAK membuat data warga otomatis.
# ============================================================

@app.route(
    "/register",
    methods=["GET", "POST"]
)
@admin_required
def register():

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        )

        konfirmasi_password = request.form.get(
            "konfirmasi_password",
            ""
        )

        if username == "":

            return render_template(
                "register.html",
                error="Username tidak boleh kosong.",
                username=username
            )

        if len(username) < 3:

            return render_template(
                "register.html",
                error="Username minimal 3 karakter.",
                username=username
            )

        if password == "":

            return render_template(
                "register.html",
                error="Password tidak boleh kosong.",
                username=username
            )

        if len(password) < 6:

            return render_template(
                "register.html",
                error="Password minimal 6 karakter.",
                username=username
            )

        if password != konfirmasi_password:

            return render_template(
                "register.html",
                error="Konfirmasi password tidak sama.",
                username=username
            )

        conn = get_db()

        existing = conn.execute("""
            SELECT id
            FROM users
            WHERE username = ?
        """, (
            username,
        )).fetchone()

        if existing:

            conn.close()

            return render_template(
                "register.html",
                error="Username sudah digunakan.",
                username=username
            )

        password_hash = generate_password_hash(
            password
        )

        conn.execute("""
            INSERT INTO users
            (
                username,
                password,
                role
            )
            VALUES (?, ?, ?)
        """, (
            username,
            password_hash,
            "user"
        ))

        conn.commit()
        conn.close()

        return redirect(
            "/admin#akun-warga"
        )

    return render_template(
        "register.html"
    )


# ============================================================
# HAPUS AKUN USER
# ============================================================

@app.route(
    "/admin/user/hapus/<int:id>"
)
@admin_required
def hapus_user(id):

    conn = get_db()

    user_data = conn.execute("""
        SELECT id
        FROM users
        WHERE id = ?
        AND role = 'user'
    """, (
        id,
    )).fetchone()

    if user_data is None:

        conn.close()

        return "Akun user tidak ditemukan."

    conn.execute("""
        DELETE FROM warga
        WHERE user_id = ?
    """, (
        id,
    ))

    conn.execute("""
        DELETE FROM pengajuan_surat
        WHERE user_id = ?
    """, (
        id,
    ))

    conn.execute("""
        DELETE FROM users
        WHERE id = ?
        AND role = 'user'
    """, (
        id,
    ))

    conn.commit()
    conn.close()

    return redirect(
        "/admin#akun-warga"
    )


# ============================================================
# PENGATURAN ADMIN
# ============================================================

@app.route(
    "/admin/pengaturan",
    methods=["GET", "POST"]
)
@admin_required
def pengaturan_admin():

    current = get_admin()

    conn = get_db()

    admin = conn.execute("""
        SELECT *
        FROM users
        WHERE id = ?
        AND role = 'admin'
    """, (
        current["user_id"],
    )).fetchone()

    if admin is None:

        conn.close()

        return "Akun admin tidak ditemukan."

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        )

        konfirmasi_password = request.form.get(
            "konfirmasi_password",
            ""
        )

        if username == "":

            conn.close()

            return render_template(
                "pengaturan.html",
                admin=admin,
                error="Username tidak boleh kosong."
            )

        if len(username) < 3:

            conn.close()

            return render_template(
                "pengaturan.html",
                admin=admin,
                error="Username minimal 3 karakter."
            )

        if password == "":

            conn.close()

            return render_template(
                "pengaturan.html",
                admin=admin,
                error="Password baru harus diisi."
            )

        if len(password) < 6:

            conn.close()

            return render_template(
                "pengaturan.html",
                admin=admin,
                error="Password minimal 6 karakter."
            )

        if password != konfirmasi_password:

            conn.close()

            return render_template(
                "pengaturan.html",
                admin=admin,
                error="Konfirmasi password tidak sama."
            )

        existing = conn.execute("""
            SELECT id
            FROM users
            WHERE username = ?
            AND id != ?
        """, (
            username,
            current["user_id"]
        )).fetchone()

        if existing:

            conn.close()

            return render_template(
                "pengaturan.html",
                admin=admin,
                error="Username sudah digunakan."
            )

        password_hash = generate_password_hash(
            password
        )

        conn.execute("""
            UPDATE users
            SET
                username = ?,
                password = ?
            WHERE id = ?
            AND role = 'admin'
        """, (
            username,
            password_hash,
            current["user_id"]
        ))

        conn.commit()
        conn.close()

        token = create_auth_token(
            current["user_id"],
            username,
            "admin"
        )

        set_session({
            "user_id": current["user_id"],
            "username": username,
            "role": "admin"
        })

        response = make_response(
            redirect("/admin")
        )

        response.set_cookie(
            ADMIN_COOKIE,
            token,
            max_age=COOKIE_MAX_AGE,
            httponly=True,
            samesite="Lax"
        )

        return response

    conn.close()

    return render_template(
        "pengaturan.html",
        admin=admin
    )


# ============================================================
# TAMBAH DATA WARGA
# ============================================================

@app.route(
    "/tambah",
    methods=["POST"]
)
@user_required
def tambah():

    current = get_user()

    nama = request.form.get(
        "nama",
        ""
    ).strip()

    nik = request.form.get(
        "nik",
        ""
    ).strip()

    alamat = request.form.get(
        "alamat",
        ""
    ).strip()

    if nama == "" or nik == "" or alamat == "":

        return "Semua data harus diisi."

    if not nik.isdigit() or len(nik) != 16:

        return "NIK harus terdiri dari 16 digit."

    conn = get_db()

    conn.execute("""
        INSERT INTO warga
        (
            nama,
            nik,
            alamat,
            user_id
        )
        VALUES (?, ?, ?, ?)
    """, (
        nama,
        nik,
        alamat,
        current["user_id"]
    ))

    conn.commit()
    conn.close()

    return redirect(
        "/user#data-warga"
    )


# ============================================================
# EDIT DATA WARGA
# ============================================================

@app.route(
    "/edit/<int:id>",
    methods=["GET", "POST"]
)
@login_required
def edit(id):

    current = get_current_user()

    conn = get_db()

    data = conn.execute("""
        SELECT *
        FROM warga
        WHERE id = ?
    """, (
        id,
    )).fetchone()

    if data is None:

        conn.close()

        return "Data warga tidak ditemukan."

    if current["role"] == "user":

        if data["user_id"] != current["user_id"]:

            conn.close()

            return "Anda tidak memiliki izin."

    if request.method == "POST":

        nama = request.form.get(
            "nama",
            ""
        ).strip()

        nik = request.form.get(
            "nik",
            ""
        ).strip()

        alamat = request.form.get(
            "alamat",
            ""
        ).strip()

        if nama == "" or nik == "" or alamat == "":

            conn.close()

            return "Semua data harus diisi."

        if not nik.isdigit() or len(nik) != 16:

            conn.close()

            return "NIK harus terdiri dari 16 digit."

        if current["role"] == "user":

            conn.execute("""
                UPDATE warga
                SET
                    nama = ?,
                    nik = ?,
                    alamat = ?
                WHERE id = ?
                AND user_id = ?
            """, (
                nama,
                nik,
                alamat,
                id,
                current["user_id"]
            ))

        else:

            conn.execute("""
                UPDATE warga
                SET
                    nama = ?,
                    nik = ?,
                    alamat = ?
                WHERE id = ?
            """, (
                nama,
                nik,
                alamat,
                id
            ))

        conn.commit()
        conn.close()

        if current["role"] == "admin":

            return redirect(
                "/admin#data-warga"
            )

        return redirect(
            "/user#data-warga"
        )

    conn.close()

    return render_template(
        "edit.html",
        data=data
    )


# ============================================================
# HAPUS WARGA
# ============================================================

@app.route(
    "/hapus/<int:id>"
)
@admin_required
def hapus(id):

    conn = get_db()

    conn.execute("""
        DELETE FROM warga
        WHERE id = ?
    """, (
        id,
    ))

    conn.commit()
    conn.close()

    return redirect(
        "/admin#data-warga"
    )


# ============================================================
# TAMBAH PENGUMUMAN
# ============================================================

@app.route(
    "/pengumuman/tambah",
    methods=["GET", "POST"]
)
@admin_required
def tambah_pengumuman():

    if request.method == "POST":

        judul = request.form.get(
            "judul",
            ""
        ).strip()

        isi = request.form.get(
            "isi",
            ""
        ).strip()

        if judul == "" or isi == "":

            return render_template(
                "pengumuman_form.html",
                error="Judul dan isi harus diisi.",
                mode="tambah"
            )

        tanggal = datetime.now().strftime(
            "%d-%m-%Y %H:%M"
        )

        conn = get_db()

        conn.execute("""
            INSERT INTO pengumuman
            (
                judul,
                isi,
                tanggal
            )
            VALUES (?, ?, ?)
        """, (
            judul,
            isi,
            tanggal
        ))

        conn.commit()
        conn.close()

        return redirect(
            "/admin#pengumuman"
        )

    return render_template(
        "pengumuman_form.html",
        mode="tambah"
    )


# ============================================================
# EDIT PENGUMUMAN
# ============================================================

@app.route(
    "/pengumuman/edit/<int:id>",
    methods=["GET", "POST"]
)
@admin_required
def edit_pengumuman(id):

    conn = get_db()

    data = conn.execute("""
        SELECT *
        FROM pengumuman
        WHERE id = ?
    """, (
        id,
    )).fetchone()

    if data is None:

        conn.close()

        return "Pengumuman tidak ditemukan."

    if request.method == "POST":

        judul = request.form.get(
            "judul",
            ""
        ).strip()

        isi = request.form.get(
            "isi",
            ""
        ).strip()

        if judul == "" or isi == "":

            conn.close()

            return render_template(
                "pengumuman_form.html",
                error="Judul dan isi harus diisi.",
                mode="edit",
                data=data
            )

        conn.execute("""
            UPDATE pengumuman
            SET
                judul = ?,
                isi = ?
            WHERE id = ?
        """, (
            judul,
            isi,
            id
        ))

        conn.commit()
        conn.close()

        return redirect(
            "/admin#pengumuman"
        )

    conn.close()

    return render_template(
        "pengumuman_form.html",
        mode="edit",
        data=data
    )


# ============================================================
# HAPUS PENGUMUMAN
# ============================================================

@app.route(
    "/pengumuman/hapus/<int:id>"
)
@admin_required
def hapus_pengumuman(id):

    conn = get_db()

    conn.execute("""
        DELETE FROM pengumuman
        WHERE id = ?
    """, (
        id,
    ))

    conn.commit()
    conn.close()

    return redirect(
        "/admin#pengumuman"
    )


# ============================================================
# LIHAT PENGUMUMAN
# ============================================================

@app.route(
    "/pengumuman"
)
@login_required
def lihat_pengumuman():

    current = get_current_user()

    conn = get_db()

    pengumuman = conn.execute("""
        SELECT *
        FROM pengumuman
        ORDER BY id DESC
    """).fetchall()

    conn.close()

    return render_template(
        "pengumuman.html",
        pengumuman=pengumuman,
        username=current["username"]
    )


# ============================================================
# SURAT ONLINE - USER
# ============================================================

@app.route(
    "/surat"
)
@user_required
def surat():

    current = get_user()

    conn = get_db()

    pengajuan = conn.execute("""
        SELECT *
        FROM pengajuan_surat
        WHERE user_id = ?
        ORDER BY id DESC
    """, (
        current["user_id"],
    )).fetchall()

    conn.close()

    return render_template(
        "surat.html",
        pengajuan=pengajuan,
        username=current["username"]
    )


# ============================================================
# AJUKAN SURAT
# ============================================================

@app.route(
    "/surat/ajukan",
    methods=["POST"]
)
@user_required
def ajukan_surat():

    current = get_user()

    jenis_surat = request.form.get(
        "jenis_surat",
        ""
    ).strip()

    keperluan = request.form.get(
        "keperluan",
        ""
    ).strip()

    keterangan = request.form.get(
        "keterangan",
        ""
    ).strip()

    if jenis_surat == "" or keperluan == "":

        return "Jenis surat dan keperluan harus diisi."

    tanggal = datetime.now().strftime(
        "%d-%m-%Y %H:%M"
    )

    conn = get_db()

    warga = conn.execute("""
        SELECT id
        FROM warga
        WHERE user_id = ?
        LIMIT 1
    """, (
        current["user_id"],
    )).fetchone()

    if warga is None:

        conn.close()

        return "Silakan isi Data Warga terlebih dahulu."

    conn.execute("""
        INSERT INTO pengajuan_surat
        (
            user_id,
            jenis_surat,
            keperluan,
            keterangan,
            tanggal_pengajuan,
            status
        )
        VALUES (?, ?, ?, ?, ?, ?)
    """, (
        current["user_id"],
        jenis_surat,
        keperluan,
        keterangan,
        tanggal,
        "Menunggu"
    ))

    conn.commit()
    conn.close()

    return redirect("/surat")


# ============================================================
# DETAIL SURAT ADMIN
# ============================================================

@app.route(
    "/admin/surat/<int:id>"
)
@admin_required
def detail_surat(id):

    conn = get_db()

    data = conn.execute("""
        SELECT
            pengajuan_surat.*,
            users.username,
            warga.nama,
            warga.nik,
            warga.alamat
        FROM pengajuan_surat

        LEFT JOIN users
        ON pengajuan_surat.user_id = users.id

        LEFT JOIN warga
        ON pengajuan_surat.user_id = warga.user_id

        WHERE pengajuan_surat.id = ?
    """, (
        id,
    )).fetchone()

    conn.close()

    if data is None:

        return "Pengajuan tidak ditemukan."

    return render_template(
        "detail_surat.html",
        data=data
    )


# ============================================================
# SETUJUI SURAT
# ============================================================

@app.route(
    "/admin/surat/<int:id>/setujui",
    methods=["POST"]
)
@admin_required
def setujui_surat(id):

    nomor_surat = request.form.get(
        "nomor_surat",
        ""
    ).strip()

    if nomor_surat == "":

        return "Nomor surat harus diisi."

    conn = get_db()

    data = conn.execute("""
        SELECT id
        FROM pengajuan_surat
        WHERE id = ?
    """, (
        id,
    )).fetchone()

    if data is None:

        conn.close()

        return "Pengajuan tidak ditemukan."

    conn.execute("""
        UPDATE pengajuan_surat
        SET
            status = 'Disetujui',
            nomor_surat = ?,
            alasan_penolakan = NULL
        WHERE id = ?
    """, (
        nomor_surat,
        id
    ))

    conn.commit()
    conn.close()

    # Tetap kembali ke ADMIN
    return redirect(
        "/admin#surat-online"
    )


# ============================================================
# TOLAK SURAT
# ============================================================

@app.route(
    "/admin/surat/<int:id>/tolak",
    methods=["POST"]
)
@admin_required
def tolak_surat(id):

    alasan = request.form.get(
        "alasan_penolakan",
        ""
    ).strip()

    if alasan == "":

        return "Alasan penolakan harus diisi."

    conn = get_db()

    data = conn.execute("""
        SELECT id
        FROM pengajuan_surat
        WHERE id = ?
    """, (
        id,
    )).fetchone()

    if data is None:

        conn.close()

        return "Pengajuan tidak ditemukan."

    conn.execute("""
        UPDATE pengajuan_surat
        SET
            status = 'Ditolak',
            alasan_penolakan = ?,
            nomor_surat = NULL
        WHERE id = ?
    """, (
        alasan,
        id
    ))

    conn.commit()
    conn.close()

    # ========================================================
    # PENTING:
    # Tetap kembali ke ADMIN.
    #
    # Tidak redirect ke /user.
    # ========================================================

    return redirect(
        "/admin#surat-online"
    )


# ============================================================
# BUAT SURAT WORD
# ============================================================

def buat_surat_word(data):

    document = Document()

    section = document.sections[0]

    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

    style = document.styles["Normal"]

    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    # --------------------------------------------------------
    # KOP
    # --------------------------------------------------------

    p = document.add_paragraph()

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run(
        "PEMERINTAH KOTA BANJARMASIN\n"
        "KECAMATAN .........................\n"
        "KELURAHAN .........................\n"
        "RT ........ RW ........"
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)

    p = document.add_paragraph()

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run(
        "____________________________________________________________"
    )

    run.bold = True

    # --------------------------------------------------------
    # JUDUL
    # --------------------------------------------------------

    p = document.add_paragraph()

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    jenis = data["jenis_surat"].upper()

    run = p.add_run(jenis)

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)

    # --------------------------------------------------------
    # NOMOR
    # --------------------------------------------------------

    p = document.add_paragraph()

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run(
        f"Nomor: {data['nomor_surat'] or '-'}"
    )

    run.bold = True

    document.add_paragraph(
        "Yang bertanda tangan di bawah ini, "
        "Ketua RT ........ RW ........, "
        "menerangkan bahwa:"
    )

    # --------------------------------------------------------
    # DATA WARGA
    # --------------------------------------------------------

    table = document.add_table(
        rows=3,
        cols=2
    )

    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    table.autofit = False

    table.columns[0].width = Cm(3)
    table.columns[1].width = Cm(12)

    data_table = [
        (
            "Nama",
            data["nama"] or "-"
        ),
        (
            "NIK",
            data["nik"] or "-"
        ),
        (
            "Alamat",
            data["alamat"] or "-"
        )
    ]

    for index, item in enumerate(data_table):

        cell1 = table.cell(index, 0)
        cell2 = table.cell(index, 1)

        cell1.text = item[0]
        cell2.text = ": " + item[1]

        cell1.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

        cell2.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

        for paragraph in cell1.paragraphs:

            for run in paragraph.runs:

                run.font.name = "Times New Roman"
                run.font.size = Pt(12)

        for paragraph in cell2.paragraphs:

            for run in paragraph.runs:

                run.font.name = "Times New Roman"
                run.font.size = Pt(12)

    document.add_paragraph("")

    # --------------------------------------------------------
    # ISI
    # --------------------------------------------------------

    jenis_lower = data["jenis_surat"].lower()

    if "domisili" in jenis_lower:

        isi = (
            "Adalah benar warga yang berdomisili "
            "di wilayah RT ........ RW ........ "
            "Kelurahan ......................... "
            "Kecamatan ......................... "
            "Kota Banjarmasin."
        )

    elif "usaha" in jenis_lower:

        isi = (
            "Adalah benar warga yang berdomisili "
            "di wilayah RT ........ RW ........ "
            "dan berdasarkan keterangan yang ada, "
            "yang bersangkutan memiliki kegiatan usaha."
        )

    elif "pengantar" in jenis_lower:

        isi = (
            "Adalah benar warga yang berdomisili "
            "di wilayah RT ........ RW ........ "
            "dan surat ini dibuat sebagai surat pengantar "
            "untuk keperluan administrasi."
        )

    else:

        isi = (
            "Adalah benar warga yang berdomisili "
            "di wilayah RT ........ RW ........ "
            "dan surat ini dibuat untuk keperluan "
            "administrasi yang bersangkutan."
        )

    document.add_paragraph(isi)

    document.add_paragraph(
        "Surat keterangan ini dibuat untuk keperluan: "
        f"{data['keperluan']}."
    )

    if data["keterangan"]:

        document.add_paragraph(
            "Keterangan tambahan: "
            f"{data['keterangan']}."
        )

    document.add_paragraph(
        "Demikian surat keterangan ini dibuat "
        "dengan sebenarnya untuk dapat dipergunakan "
        "sebagaimana mestinya."
    )

    tanggal = datetime.now().strftime(
        "%d %B %Y"
    )

    p = document.add_paragraph()

    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    p.add_run(
        f"Banjarmasin, {tanggal}\n\n"
        "Ketua RT ........ RW ........\n\n\n\n"
        "(________________________)"
    )

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output


# ============================================================
# DOWNLOAD SURAT
# ============================================================

@app.route(
    "/surat/download/<int:id>"
)
@login_required
def download_surat(id):

    current = get_current_user()

    conn = get_db()

    data = conn.execute("""
        SELECT
            pengajuan_surat.*,
            users.username,
            warga.nama,
            warga.nik,
            warga.alamat
        FROM pengajuan_surat

        LEFT JOIN users
        ON pengajuan_surat.user_id = users.id

        LEFT JOIN warga
        ON pengajuan_surat.user_id = warga.user_id

        WHERE pengajuan_surat.id = ?
    """, (
        id,
    )).fetchone()

    conn.close()

    if data is None:

        return "Surat tidak ditemukan."

    if current["role"] == "user":

        if data["user_id"] != current["user_id"]:

            return "Anda tidak memiliki izin."

    if data["status"] != "Disetujui":

        return "Surat belum disetujui oleh Admin."

    output = buat_surat_word(data)

    nama_file = (
        data["jenis_surat"]
        .replace(" ", "_")
        .replace("/", "_")
        .replace("\\", "_")
    )

    nama_warga = (
        data["nama"]
        .replace(" ", "_")
        .replace("/", "_")
        .replace("\\", "_")
        if data["nama"]
        else "warga"
    )

    filename = (
        f"{nama_file}_{nama_warga}.docx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )
    )


# ============================================================
# EXPORT EXCEL
# ============================================================

@app.route(
    "/admin/export-excel"
)
@admin_required
def export_excel():

    conn = get_db()

    data = conn.execute("""
        SELECT
            pengajuan_surat.id,
            warga.nama,
            warga.nik,
            warga.alamat,
            pengajuan_surat.jenis_surat,
            pengajuan_surat.keperluan,
            pengajuan_surat.keterangan,
            pengajuan_surat.tanggal_pengajuan,
            pengajuan_surat.status,
            pengajuan_surat.alasan_penolakan,
            pengajuan_surat.nomor_surat
        FROM pengajuan_surat

        LEFT JOIN warga
        ON pengajuan_surat.user_id = warga.user_id

        ORDER BY pengajuan_surat.id DESC
    """).fetchall()

    conn.close()

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Pengajuan Surat"

    headers = [
        "No",
        "Nama",
        "NIK",
        "Alamat",
        "Jenis Surat",
        "Keperluan",
        "Keterangan",
        "Tanggal Pengajuan",
        "Status",
        "Alasan Penolakan",
        "Nomor Surat"
    ]

    sheet.append(headers)

    for item in data:

        sheet.append([
            item["id"],
            item["nama"],
            item["nik"],
            item["alamat"],
            item["jenis_surat"],
            item["keperluan"],
            item["keterangan"],
            item["tanggal_pengajuan"],
            item["status"],
            item["alasan_penolakan"],
            item["nomor_surat"]
        ])

    widths = {
        "A": 8,
        "B": 25,
        "C": 20,
        "D": 35,
        "E": 25,
        "F": 30,
        "G": 35,
        "H": 20,
        "I": 15,
        "J": 30,
        "K": 25
    }

    for column, width in widths.items():

        sheet.column_dimensions[
            column
        ].width = width

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="pengajuan_surat.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


# ============================================================
# LOGOUT
# ============================================================

@app.route("/logout")
def logout():

    # Logout dibuat terpisah agar logout dari satu tab
    # tidak ikut mengeluarkan akun pada tab lainnya.
    # Jika parameter role tidak diberikan, gunakan halaman asal.
    role = request.args.get("role", "").strip().lower()
    referrer = request.referrer or ""

    if role not in ("admin", "user"):
        if "/admin" in referrer:
            role = "admin"
        elif "/user" in referrer or "/surat" in referrer:
            role = "user"

    response = make_response(redirect("/"))

    if role == "admin":
        response.delete_cookie(ADMIN_COOKIE)
        session.clear()
    elif role == "user":
        response.delete_cookie(USER_COOKIE)
        session.clear()
    else:
        response.delete_cookie(ADMIN_COOKIE)
        response.delete_cookie(USER_COOKIE)
        session.clear()

    return response


# ============================================================
# START
# ============================================================

if __name__ == "__main__":

    init_db()

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )

